#!/usr/bin/env python3
"""
extract_plaintext.py — Extract plain text from a DOCX or text-based PDF for the
parse-verification step.

What the parser sees is approximately what this script extracts. If the output is
scrambled, the ATS parser will scramble it too — fix the source before submitting.

Usage:
    python extract_plaintext.py firstname-lastname-resume.docx
    python extract_plaintext.py firstname-lastname-resume.pdf

Writes a sibling .txt file next to the input (e.g.,
firstname-lastname-parse-check.txt).

Requires:
    python-docx for .docx (pip install python-docx)
    pypdf for .pdf      (pip install pypdf)
"""

import sys
from pathlib import Path


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        print("Error: python-docx is required for .docx. Install: pip install python-docx",
              file=sys.stderr)
        sys.exit(1)
    doc = Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.rstrip()
        if text:
            lines.append(text)
        else:
            # Preserve blank line as section break
            if lines and lines[-1] != "":
                lines.append("")
    # Tables (if any) — flag with warning since they shouldn't exist in a
    # format-compliant resume
    if doc.tables:
        lines.append("")
        lines.append("[WARNING] Document contains tables — these often scramble in "
                     "ATS parsers. Convert to single-column paragraph layout.")
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip("| "):
                    lines.append(row_text)
    return "\n".join(lines).strip() + "\n"


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        print("Error: pypdf is required for .pdf. Install: pip install pypdf",
              file=sys.stderr)
        sys.exit(1)
    reader = PdfReader(str(path))
    chunks: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        chunks.append(text.strip())
    return "\n\n".join(chunks).strip() + "\n"


def parse_check_path(input_path: Path) -> Path:
    """Build the output path: replace -resume.<ext> with -parse-check.txt if the
    convention is in use, else append -parse-check.txt to the stem."""
    stem = input_path.stem
    if stem.endswith("-resume"):
        stem = stem[: -len("-resume")]
    return input_path.with_name(f"{stem}-parse-check.txt")


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__, file=sys.stderr)
        return 2
    input_path = Path(argv[1])
    if not input_path.exists():
        print(f"Error: {input_path} not found", file=sys.stderr)
        return 1
    suffix = input_path.suffix.lower()
    if suffix == ".docx":
        text = extract_docx(input_path)
    elif suffix == ".pdf":
        text = extract_pdf(input_path)
    else:
        print(f"Error: unsupported format {suffix}. Use .docx or .pdf.", file=sys.stderr)
        return 1
    output_path = Path(argv[2]) if len(argv) >= 3 else parse_check_path(input_path)
    output_path.write_text(text, encoding="utf-8")
    print(f"Wrote {output_path}")
    print()
    print("Review checklist (see references/parse-verification.md for full diagnostics):")
    print("  [ ] Contact info on lines 1-5")
    print("  [ ] Section headings present as their own lines")
    print("  [ ] Job titles, employers, and dates on adjacent lines per role")
    print("  [ ] Dates in 'Mon YYYY – Mon YYYY' format")
    print("  [ ] Bullets begin with '•' character (not '?' or garbage)")
    print("  [ ] Skills section reads as a flat list, not scrambled")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
